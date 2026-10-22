import io

import pandas as pd
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# cspell:ignore pagesizes


def make_xlsx(summary, issues):
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        pd.DataFrame(
            [
                {"Metric": k, "Value": str(v)}
                for k, v in summary.items()
                if k != "issue_counts"
            ]
        ).to_excel(w, index=False, sheet_name="Summary")
        pd.DataFrame(
            [
                {"Issue type": k, "Count": v}
                for k, v in summary.get("issue_counts", {}).items()
            ]
        ).to_excel(w, index=False, sheet_name="Issue counts")
        pd.DataFrame(issues).to_excel(w, index=False, sheet_name="Detailed issues")
    out.seek(0)
    return out.getvalue()


def make_pdf(summary, issues):
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), rightMargin=24, leftMargin=24)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("SheetWise — File Comparison Report", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            f"Status: {summary['status']} | Total findings: {summary['total_issues']}",
            styles["Heading2"],
        ),
        Paragraph(
            f"Reference: {summary['reference_rows']} rows × {summary['reference_columns']} columns; Comparison: {summary['comparison_rows']} rows × {summary['comparison_columns']} columns",
            styles["BodyText"],
        ),
        Spacer(1, 12),
    ]
    rows = [
        ["Issue type", "Record key", "Column", "Reference value", "Comparison value"]
    ]
    for x in issues[:2000]:
        rows.append(
            [
                str(x.get("issue_type", "")),
                str(x.get("record_key", "")),
                str(x.get("column", "")),
                str(x.get("reference_value", ""))[:100],
                str(x.get("comparison_value", ""))[:100],
            ]
        )
    t = Table(rows, repeatRows=1, colWidths=[100, 75, 90, 210, 210])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f1f5f9")],
                ),
            ]
        )
    )
    story.append(t)
    if len(issues) > 2000:
        story.append(
            Paragraph(
                "PDF shows the first 2,000 findings. Use Excel for the full list.",
                styles["Italic"],
            )
        )
    doc.build(story)
    out.seek(0)
    return out.getvalue()


def make_docx(summary, issues):
    d = Document()
    d.add_heading("SheetWise — File Comparison Report", 0)
    d.add_paragraph(
        f"Status: {summary['status']} | Total findings: {summary['total_issues']}"
    )
    d.add_paragraph(
        f"Reference: {summary['reference_rows']} rows × {summary['reference_columns']} columns"
    )
    d.add_paragraph(
        f"Comparison: {summary['comparison_rows']} rows × {summary['comparison_columns']} columns"
    )
    d.add_heading("Issue counts", level=1)
    for k, v in summary.get("issue_counts", {}).items():
        d.add_paragraph(f"{k}: {v}", style="List Bullet")
    d.add_heading("Detailed findings", level=1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Shading Accent 1"
    heads = [
        "Issue type",
        "Record key",
        "Column",
        "Reference value",
        "Comparison value",
    ]
    for c, h in zip(t.rows[0].cells, heads):
        c.text = h
    for x in issues:
        vals = [
            x.get("issue_type", ""),
            x.get("record_key", ""),
            x.get("column", ""),
            x.get("reference_value", ""),
            x.get("comparison_value", ""),
        ]
        for c, v in zip(t.add_row().cells, vals):
            c.text = str(v)[:500]
    out = io.BytesIO()
    d.save(out)
    out.seek(0)
    return out.getvalue()
